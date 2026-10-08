#!/usr/bin/env python3
"""Resource Tracker seed import (deterministic).

Reads the Immersive Homes Resource Planner workbook and writes the standalone
Resource Tracker seed (src/rt/seed.json) plus an import report.

  python3 scripts/rt_import.py <workbook.xlsx> [src/rt/seed.json] [docs/rt-import-report.md]

Rules
- Only the editable task area of 'Task Input' (columns A-I) is read. PivotTable caches
  and example rows (e.g. "Example Project") in columns K+ are ignored.
- A task row is imported when Project or Task (or any of Start..Notes) is non-blank.
- Missing values stay null. Hrs/Week is never turned into 0.
- Text is whitespace-trimmed; project names are matched case-insensitively after trimming
  so "Morgan " and "Morgan" are one project. No other value is altered.
- IDs are stable UUIDs derived from names (uuid5), so re-running gives identical output.
- The Excel roster is imported first; Team US, Abhishek and Amin are added afterwards
  and skipped if a person with the same name already exists.
"""
import json, re, sys, uuid, datetime
from collections import OrderedDict
import openpyxl

NS = uuid.UUID('6f1c3a7e-2b2d-4b8a-9a55-1e2a0c0f7a11')
uid = lambda kind, key: str(uuid.uuid5(NS, f'{kind}:{key.lower()}'))
NOW = '2026-10-08T00:00:00Z'  # fixed for deterministic output

# Workbook block header -> Resource Tracker discipline name (order = display order after Company Management)
BLOCKS = OrderedDict([
    ('ARCHITECTURE', 'Architecture'),
    ('INTERIOR DESIGN', 'Interior Design / Kitchen / Cabinetry'),
    ('TECHNICAL / BIM', 'BIM Heroes'),
    ('CIVIL', 'Civil'),
    ('STRUCTURAL', 'Structural'),
    ('MEP', 'MEP'),
    ('MANUFACTURING', 'Manufacturing'),
    ('PROJECT MANAGEMENT', 'Project Management'),
])
DISPLAY_ORDER = ['Company Management', 'Architecture', 'Interior Design / Kitchen / Cabinetry', 'BIM Heroes', 'Civil',
                 'Structural', 'MEP', 'Team US', 'Manufacturing', 'Project Management']
ADDED_PEOPLE = [  # added after the Excel roster (name, discipline, role)
    ('Amin', 'Company Management', 'Company Manager'),
    ('Dwight', 'Team US', None), ('David W', 'Team US', None), ('JD', 'Team US', None),
    ('Angel', 'Team US', None), ('Mark Kemp', 'Team US', None), ('Karen', 'Team US', None),
    ('Abhishek', 'BIM Heroes', None),
]

def txt(v):
    if v is None: return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v
def date(v):
    if v is None or v == '': return None
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    raise ValueError(f'Unexpected date value {v!r}')
def num(v):
    if v is None or v == '': return None
    return float(v) if not float(v).is_integer() else int(v)

def main(path, out_seed, out_report):
    wb = openpyxl.load_workbook(path)
    report = OrderedDict(); notes = []

    # ── Disciplines ─────────────────────────────────────────────────────────
    disciplines = OrderedDict()
    for i, name in enumerate(DISPLAY_ORDER):
        disciplines[name] = {'id': uid('discipline', name), 'name': name, 'kind': 'management' if name == 'Company Management' else 'discipline',
                             'sort_order': i * 10, 'active': True, 'created_at': NOW, 'updated_at': NOW}

    # ── Team & Capacity ─────────────────────────────────────────────────────
    tc = wb['Team & Capacity']
    people = OrderedDict(); current = None; header = False; sort = {}
    for row in tc.iter_rows(min_row=1, max_row=tc.max_row, max_col=5):
        a = txt(row[0].value)
        if isinstance(a, str) and a.upper() in BLOCKS and all(c.value in (None, '') for c in row[1:5]):
            current = BLOCKS[a.upper()]; header = False; continue
        if a == 'Person': header = True; continue
        if current and header and a and not str(a).startswith('='):
            name, cap, role, active = a, num(row[1].value), txt(row[2].value), txt(row[3].value)
            key = name.lower()
            if key in people: notes.append(f'Duplicate person "{name}" in Team & Capacity skipped'); continue
            sort[current] = sort.get(current, 0) + 1
            people[key] = {'id': uid('person', name), 'discipline_id': disciplines[current]['id'], 'name': name, 'role': role,
                           'weekly_capacity_hours': cap, 'sort_order': sort[current] * 10,
                           'active': (active or 'Yes').lower() != 'no', 'created_at': NOW, 'updated_at': NOW}
    report['People imported from Excel'] = len(people)
    added = 0
    for name, disc, role in ADDED_PEOPLE:
        if name.lower() in people: notes.append(f'{name} already in the workbook; not added again'); continue
        sort[disc] = sort.get(disc, 0) + 1
        people[name.lower()] = {'id': uid('person', name), 'discipline_id': disciplines[disc]['id'], 'name': name, 'role': role,
                                'weekly_capacity_hours': None, 'sort_order': sort[disc] * 10, 'active': True, 'created_at': NOW, 'updated_at': NOW}
        added += 1
    report['People added (Team US, Abhishek, Amin)'] = added

    # ── Projects ────────────────────────────────────────────────────────────
    pr = wb['Projects']; projects = OrderedDict(); hdr = None
    for row in pr.iter_rows(min_row=1, max_row=pr.max_row, max_col=6):
        vals = [c.value for c in row]
        if txt(vals[0]) == 'Project Name': hdr = True; continue
        if not hdr or not txt(vals[0]): continue
        name = txt(vals[0]); key = name.lower()
        if key in projects: notes.append(f'Duplicate project "{name}" skipped'); continue
        projects[key] = {'id': uid('project', name), 'name': name, 'project_manager': txt(vals[1]), 'status': txt(vals[2]),
                         'start_date': date(vals[3]), 'target_finish': date(vals[4]), 'notes': txt(vals[5]), 'active': True,
                         'created_at': NOW, 'updated_at': NOW}
    report['Projects imported'] = len(projects)

    # ── Task Input (editable area A–I only) ─────────────────────────────────
    ti = wb['Task Input']; tasks = []; unmatched_people = []; unmatched_projects = []; normalized = set()
    for row in ti.iter_rows(min_row=1, max_row=ti.max_row, max_col=9):
        v = [c.value for c in row]
        if txt(v[1]) == 'Project' or not any(x not in (None, '') for x in v[1:9]):
            continue
        raw = v[0]; pname = None
        m = re.search(r"'Team & Capacity'!\$?A\$?(\d+)", str(raw or ''))
        if m: pname = txt(tc.cell(int(m.group(1)), 1).value)
        elif raw and not str(raw).startswith('='): pname = txt(raw)
        person = people.get((pname or '').lower())
        if not person: unmatched_people.append(f'row {row[0].row}: {pname!r}'); continue
        proj_raw = v[1] if isinstance(v[1], str) else (str(v[1]) if v[1] is not None else None)
        proj_name = txt(proj_raw); project = projects.get((proj_name or '').lower()) if proj_name else None
        if proj_name and proj_raw != proj_name: normalized.add(f'"{proj_raw}" → "{proj_name}"')
        if proj_name and not project: unmatched_projects.append(f'row {row[0].row}: {proj_name!r}'); continue
        tasks.append({'id': uid('task', f'{row[0].row}:{person["name"]}:{proj_name}:{txt(v[2])}'), 'person_id': person['id'],
                      'project_id': project['id'] if project else None, 'title': txt(v[2]), 'start_date': date(v[3]), 'due_date': date(v[4]),
                      'hours_per_week': num(v[5]), 'status': txt(v[6]), 'priority': txt(v[7]), 'notes': txt(v[8]), 'phase_key': None,  # future PM Workflow phase (PH1..PH6); never guessed
                      'sort_order': row[0].row * 10, 'active': True, 'source_row': row[0].row, 'created_at': NOW, 'updated_at': NOW})
    report['Tasks imported'] = len(tasks)
    report['Tasks with missing Start'] = sum(1 for t in tasks if not t['start_date'])
    report['Tasks with missing Due'] = sum(1 for t in tasks if not t['due_date'])
    report['Tasks with missing Hrs/Week'] = sum(1 for t in tasks if t['hours_per_week'] is None)
    report['Tasks with missing Status'] = sum(1 for t in tasks if not t['status'])
    report['Tasks with missing Priority'] = sum(1 for t in tasks if not t['priority'])
    report['Unmatched people'] = len(unmatched_people)
    report['Unmatched projects'] = len(unmatched_projects)

    seed = {'version': 1, 'source': 'Immersive_Homes_Resource_Planner_v4.xlsx', 'imported_at': NOW,
            'rt_disciplines': list(disciplines.values()), 'rt_people': list(people.values()),
            'rt_projects': list(projects.values()), 'rt_tasks': tasks}
    with open(out_seed, 'w', encoding='utf-8') as f: json.dump(seed, f, indent=1, ensure_ascii=False)

    by_person = OrderedDict()
    for t in tasks:
        n = next(p['name'] for p in people.values() if p['id'] == t['person_id']); by_person[n] = by_person.get(n, 0) + 1
    lines = ['# Resource Tracker — Excel import report', '', f'Source: `{seed["source"]}` (sheets: Projects, Team & Capacity, Task Input)', '', '```']
    lines += [f'{k + ":":<44}{v}' for k, v in report.items()]
    lines += ['```', '', '## Tasks per person', '', '| Person | Tasks |', '|---|---:|'] + [f'| {k} | {v} |' for k, v in by_person.items()]
    lines += ['', '## Notes', '']
    if normalized: lines.append(f'- Project names trimmed to match the Projects sheet: {", ".join(sorted(normalized))}')
    lines += [f'- {n}' for n in notes]
    if unmatched_people: lines.append(f'- Unmatched people: {"; ".join(unmatched_people)}')
    if unmatched_projects: lines.append(f'- Unmatched projects: {"; ".join(unmatched_projects)}')
    lines.append('- Ignored: PivotTable caches and the "Example Project" rows in Task Input columns K–Q; the Manager Timeline and System Data sheets (formulas only).')
    lines.append('- People added after the import (Team US, Abhishek, Amin) have no weekly capacity yet: the workbook gives none, so it is left blank rather than invented.')
    with open(out_report, 'w', encoding='utf-8') as f: f.write('\n'.join(lines) + '\n')
    print('\n'.join(lines[4:4 + len(report) + 2]))

if __name__ == '__main__':
    a = sys.argv[1:]
    main(a[0], a[1] if len(a) > 1 else 'src/rt/seed.json', a[2] if len(a) > 2 else 'docs/rt-import-report.md')
